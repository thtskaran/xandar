"""Bounded, local-only document extraction. Original documents are never modified."""
import sys, os, io, json, zipfile, resource, datetime, logging
from pathlib import Path
resource.setrlimit(resource.RLIMIT_AS, (512*1024*1024, 512*1024*1024))
resource.setrlimit(resource.RLIMIT_CPU, (12, 12))
sys.path.insert(0, str(Path(__file__).resolve().parent.parent/'vendor'/'office'))
logging.disable(logging.CRITICAL)
MAX_INPUT=16*1024*1024
MAX_TEXT=4*1024*1024
MAX_EXPANDED=32*1024*1024
class Limited(Exception): pass
class Invalid(Exception): pass

def empty(status,warning):
 return dict(kind='text',status=status,columns=[],rows=[],chunks=[],warnings=[warning],extractor='office-worker-v1')

def check_zip(raw,expected):
 if raw.startswith(b'\xd0\xcf\x11\xe0'):
  return None
 try: z=zipfile.ZipFile(io.BytesIO(raw))
 except zipfile.BadZipFile: raise Invalid('Not a valid Office ZIP document.')
 members=z.infolist()
 if len(members)>2048: raise Limited('Office archive has more than 2,048 members.')
 if len({m.filename for m in members})!=len(members): raise Invalid('Office archive contains duplicate member names.')
 if expected not in z.namelist(): raise Invalid('Office archive does not contain the expected document part.')
 total=0
 for m in members:
  if m.flag_bits&1: raise Invalid('Encrypted ZIP members are unsupported.')
  if m.filename.startswith('/') or '..' in m.filename.split('/') or '\\' in m.filename: raise Invalid('Invalid Office archive member name.')
  total+=m.file_size
  if m.file_size>16*1024*1024 or total>MAX_EXPANDED: raise Limited('Office expanded-byte limit exceeded (16 MiB/member, 32 MiB total).')
  if m.file_size>max(1024*1024,m.compress_size*150): raise Limited('Office member compression ratio exceeds the supported bound.')
  if m.compress_type not in [zipfile.ZIP_STORED,zipfile.ZIP_DEFLATED]: raise Invalid('Unsupported Office archive compression.')
  # Bounded read validates actual expansion and CRC; no files are extracted to disk.
  with z.open(m) as f:
   blob=f.read(m.file_size+1)
   if len(blob)!=m.file_size: raise Invalid('Office archive member size mismatch.')
  if m.filename.lower().endswith(('.xml','.rels')):
   from defusedxml.ElementTree import fromstring
   fromstring(blob,forbid_dtd=True,forbid_entities=True,forbid_external=True)
 return z

def cell_value(cell):
 value=cell.value
 if cell.data_type=='f': return '',True
 if value is None:return '',False
 if isinstance(value,(datetime.datetime,datetime.date)):return value.isoformat().split('T')[0],False
 if isinstance(value,bool):return str(value).lower(),False
 text=str(value)
 if len(text)>16000:raise Limited('Office cell exceeds 16,000 characters.')
 return text,False

def xlsx(raw):
 z=check_zip(raw,'xl/workbook.xml')
 if z is None:return empty('encrypted or legacy Office','Encrypted or legacy compound Office container: extraction unsupported; no password processing.')
 if any('vbaproject' in n.lower() for n in z.namelist()):return empty('macro-enabled unsupported','Macro-containing workbook stored without extraction. Use a plain XLSX export.')
 from openpyxl import load_workbook
 wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False,keep_vba=False)
 if len(wb.worksheets)>32:raise Limited('Workbook exceeds 32 sheets.')
 sheets=[];total_rows=0;total_cells=0;formulas=0
 for ws in wb.worksheets:
  if (ws.max_column or 0)>128 or (ws.max_row or 0)>50001:raise Limited('Sheet bounds exceed 128 columns or 50,000 records.')
  ws.reset_dimensions()
  headers=None;rows=[];row_numbers=[];header_row=None;sheet_formulas=0
  for number,cells in enumerate(ws.iter_rows(),1):
   if number>50001 or len(cells)>128:raise Limited('Sheet bounds exceed 128 columns or 50,000 records.')
   total_cells+=len(cells)
   if total_cells>1000000:raise Limited('Workbook exceeds one million cells.')
   values=[]
   for cell in cells:
    value,is_formula=cell_value(cell);values.append(value);sheet_formulas+=int(is_formula)
   while values and values[-1]=='':values.pop()
   if not any(values):continue
   if headers is None:
    headers=[];used=set();header_row=number
    for i,value in enumerate(values):
     base=value.strip() or f'Column {i+1}';name=base;n=2
     while name in used:name=f'{base} ({n})';n+=1
     headers.append(name);used.add(name)
    continue
   if len(values)>len(headers):raise Invalid(f'Sheet {ws.title}: data extends beyond the header; use a rectangular table.')
   rows.append(values+['']*(len(headers)-len(values)));row_numbers.append(number);total_rows+=1
   if total_rows>50000:raise Limited('Workbook exceeds 50,000 total records.')
  formulas+=sheet_formulas
  sheets.append(dict(kind='table',name=ws.title,columns=headers or [],rows=rows,rowNumbers=row_numbers,headerRow=header_row,sheetState=ws.sheet_state,formulaCellsExcluded=sheet_formulas,chunks=[]))
 wb.close()
 warnings=['All sheets inspected as rectangular tables; first nonempty row is the header. Dates become ISO dates. Styles, charts, drawings and external links are not extracted.']
 if formulas:warnings.append(f'{formulas} formula cells excluded; formulas are never evaluated and cached values are not trusted. Review financial exclusions.')
 return dict(kind='workbook',status='indexed',columns=[],rows=[],chunks=[],sheets=sheets,warnings=warnings,extractor='openpyxl 3.1.5',formulaCellsExcluded=formulas)

def docx(raw):
 z=check_zip(raw,'word/document.xml')
 if z is None:return empty('encrypted or legacy Office','Encrypted or legacy compound Office container: extraction unsupported; no password processing.')
 if any('vbaproject' in n.lower() for n in z.namelist()):return empty('macro-enabled unsupported','Macro-containing Word file stored without extraction.')
 from docx import Document
 from docx.text.paragraph import Paragraph
 doc=Document(io.BytesIO(raw));chunks=[];chars=0;paragraph=0;table=0;heading='Document body'
 def add(text,section,anchor):
  nonlocal chars
  if not text.strip():return
  chars+=len(text)
  if chars>MAX_TEXT:raise Limited('Document exceeds 4 million extracted characters.')
  for pos in range(0,len(text),2200):
   chunks.append(dict(text=text[pos:pos+2200],section=section,**anchor))
   if len(chunks)>20000:raise Limited('Document exceeds 20,000 extracted passages.')
 for block in doc.iter_inner_content():
  if isinstance(block,Paragraph):
   paragraph+=1
   if paragraph>20000:raise Limited('Document exceeds 20,000 paragraphs.')
   if block.style and block.style.name.startswith('Heading'):heading=block.text[:160] or heading
   add(block.text,heading,dict(paragraphStart=paragraph,paragraphEnd=paragraph))
  else:
   table+=1
   if table>200:raise Limited('Document exceeds 200 tables.')
   for r,row in enumerate(block.rows,1):
    if r>5000 or len(row.cells)>128:raise Limited('Word table exceeds 5,000 rows or 128 columns.')
    for c,cell in enumerate(row.cells,1):add(cell.text,f'Table {table}',dict(table=table,tableRow=r,tableColumn=c))
 return dict(kind='text',status='indexed' if chunks else 'no extractable text',columns=[],rows=[],chunks=chunks,warnings=['Extracts body paragraphs and top-level table cells only. Headers, footers, text boxes, nested tables, tracked revisions, images and embedded objects are not included; no OCR. Paragraph/table anchors are structural, not page numbers.'],extractor='python-docx 1.2.0')

def pdf(raw):
 if not raw.startswith(b'%PDF-'):raise Invalid('Not a valid PDF signature.')
 from pypdf import PdfReader, filters
 for setting in ['MAX_DECLARED_STREAM_LENGTH','MAX_ARRAY_BASED_STREAM_OUTPUT_LENGTH','JBIG2_MAX_OUTPUT_LENGTH','LZW_MAX_OUTPUT_LENGTH','RUN_LENGTH_MAX_OUTPUT_LENGTH','ZLIB_MAX_OUTPUT_LENGTH','FLATE_MAX_BUFFER_SIZE']:
  setattr(filters,setting,8*1024*1024)
 reader=PdfReader(io.BytesIO(raw),strict=True)
 if reader.is_encrypted:return empty('encrypted PDF','Encrypted PDF stored without extraction. Supply an unencrypted copy; no passwords are attempted.')
 if len(reader.pages)>100:raise Limited('PDF exceeds 100 pages.')
 chunks=[];chars=0;empty_pages=[]
 for number,page in enumerate(reader.pages,1):
  text=page.extract_text() or '';chars+=len(text)
  if chars>MAX_TEXT:raise Limited('PDF exceeds 4 million extracted characters.')
  if not text.strip():empty_pages.append(number)
  for pos in range(0,len(text),2200):chunks.append(dict(text=text[pos:pos+2200],section=f'Page {number}',page=number))
 warnings=['PDF text reading order may differ from visual layout. Tables are text only. Images, attachments, scripts and links are not executed or extracted. No OCR.']
 if empty_pages:warnings.append(f'{len(empty_pages)} of {len(reader.pages)} pages have no extractable text; they may be scanned, image-only or blank. No OCR was run.')
 return dict(kind='text',status='indexed' if chunks else 'no extractable text',columns=[],rows=[],chunks=chunks,warnings=warnings,pageCount=len(reader.pages),emptyTextPages=empty_pages,extractor='pypdf 6.16.2')

def main():
 try:
  raw=sys.stdin.buffer.read(MAX_INPUT+1)
  if len(raw)>MAX_INPUT:raise Limited('Document exceeds 16 MiB.')
  ext=sys.argv[1]
  result={'xlsx':xlsx,'docx':docx,'pdf':pdf}[ext](raw)
  print(json.dumps({'ok':True,'data':result},ensure_ascii=False,separators=(',',':')))
 except ImportError:print(json.dumps({'ok':False,'code':'dependency_missing','message':'Document extraction dependencies are unavailable; run the documented local setup.'}))
 except Limited as e:print(json.dumps({'ok':False,'code':'limit','message':str(e)}))
 except (MemoryError,RecursionError):print(json.dumps({'ok':False,'code':'limit','message':'Document complexity exceeds worker memory or nesting limits.'}))
 except Invalid as e:print(json.dumps({'ok':False,'code':'invalid','message':str(e)}))
 except Exception:print(json.dumps({'ok':False,'code':'invalid','message':'Document is malformed or uses unsupported format features.'}))
if __name__=='__main__':main()
