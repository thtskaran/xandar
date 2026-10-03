from pathlib import Path
import sys,zipfile
root=Path(__file__).resolve().parent.parent
sys.path.insert(0,str(root/'vendor'/'office'))
from openpyxl import Workbook
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
out=root/'test'/'office-fixtures';out.mkdir(exist_ok=True)
w=Workbook();s=w.active;s.title='October orders';s.append(['order_id','order_date','customer_id','total','currency','status']);s.append(['A','2026-10-01','C1',100,'INR','completed']);s.append(['B','2026-10-02','C2',250,'INR','paid']);s.append(['Formula','2026-10-02','C2','=100+1','INR','paid']);s2=w.create_sheet('November orders');s2.append(['order_id','order_date','customer_id','total','currency','status']);s2.append(['C','2026-11-01','C1',75,'INR','completed']);s2.append(['D','2026-11-02','C2',999,'INR','cancelled']);w.save(out/'orders-two-sheets.xlsx')
w=Workbook();w.active.cell(row=50002,column=1,value='out of bounds');w.save(out/'oversized-dimensions.xlsx')
d=Document();d.add_heading('Cold chain policy',0);d.add_paragraph('JuiceShop fictional freshness checks happen before dispatch.');t=d.add_table(rows=2,cols=2);t.cell(0,0).text='Stage';t.cell(0,1).text='Owner';t.cell(1,0).text='Delivery';t.cell(1,1).text='Operations';d.save(out/'policy.docx')
writer=PdfWriter();page=writer.add_blank_page(width=612,height=792);font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')});page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})});stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 72 720 Td (JuiceShop fictional beverage delivery policy.) Tj ET');page[NameObject('/Contents')]=writer._add_object(stream);writer.add_blank_page(width=612,height=792);writer.write(out/'text-and-blank.pdf');writer.encrypt('fixture-password');writer.write(out/'encrypted.pdf')
w=PdfWriter();w.add_blank_page(width=612,height=792);w.write(out/'blank.pdf')
w=PdfWriter()
for i in range(101):w.add_blank_page(width=612,height=792)
w.write(out/'over-100-pages.pdf')
with zipfile.ZipFile(out/'oversized-member.docx','w',zipfile.ZIP_DEFLATED) as z:z.writestr('word/document.xml',b'a'*(17*1024*1024))
(out/'invalid.docx').write_bytes(b'not an Office document')
print('Created 9 real and invalid/bounded Office/PDF fixtures.')
